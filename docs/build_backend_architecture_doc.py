from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path(__file__).with_name("Duit_Cards_Backend_Architecture_Phase_1.docx")

BLACK = RGBColor(0, 0, 0)
BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
MUTED = RGBColor(89, 89, 89)
LIGHT_GRAY = "F2F4F7"
CALLOUT = "F4F6F9"
BORDER = "D9E2F3"


def set_run_font(run, name="Calibri", size=11, color=BLACK, bold=None, italic=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa, indent_dxa=120):
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")

    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = tbl.tblGrid
    if grid is None:
        grid = OxmlElement("w:tblGrid")
        tbl.append(grid)
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            set_cell_width(cell, widths_dxa[idx])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_table_borders(table, color="D9D9D9", size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        node = borders.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def set_paragraph_spacing(paragraph, before=0, after=6, line=1.10):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line


def add_para(doc, text="", size=11, bold=False, color=BLACK, italic=False, after=6, before=0, align=None):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    set_paragraph_spacing(p, before=before, after=after)
    r = p.add_run(text)
    set_run_font(r, size=size, color=color, bold=bold, italic=italic)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    if level == 1:
        size, color, before, after = 16, BLUE, 16, 8
    elif level == 2:
        size, color, before, after = 13, BLUE, 12, 6
    else:
        size, color, before, after = 12, DARK_BLUE, 8, 4
    set_paragraph_spacing(p, before=before, after=after)
    r = p.add_run(text)
    set_run_font(r, size=size, color=color, bold=True)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    set_paragraph_spacing(p, after=4, line=1.167)
    r = p.add_run(text)
    set_run_font(r, size=11, color=BLACK)
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    set_paragraph_spacing(p, after=4, line=1.167)
    r = p.add_run(text)
    set_run_font(r, size=11, color=BLACK)
    return p


def add_callout(doc, title, body):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    set_table_geometry(table, [9360], indent_dxa=120)
    set_table_borders(table, color=BORDER, size="6")
    cell = table.cell(0, 0)
    shade_cell(cell, CALLOUT)
    p = cell.paragraphs[0]
    set_paragraph_spacing(p, after=4)
    r = p.add_run(title)
    set_run_font(r, size=11, color=DARK_BLUE, bold=True)
    p2 = cell.add_paragraph()
    set_paragraph_spacing(p2, after=0)
    r2 = p2.add_run(body)
    set_run_font(r2, size=10.5, color=BLACK)
    doc.add_paragraph()
    return table


def add_table(doc, headers, rows, widths_dxa):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.style = "Table Grid"
    set_table_geometry(table, widths_dxa)
    set_table_borders(table)

    header_cells = table.rows[0].cells
    tr_pr = table.rows[0]._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)

    for i, header in enumerate(headers):
        shade_cell(header_cells[i], LIGHT_GRAY)
        p = header_cells[i].paragraphs[0]
        set_paragraph_spacing(p, after=0)
        r = p.add_run(header)
        set_run_font(r, size=10.5, color=BLACK, bold=True)

    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            p = cells[i].paragraphs[0]
            set_paragraph_spacing(p, after=0, line=1.10)
            r = p.add_run(value)
            set_run_font(r, size=10, color=BLACK)

    set_table_geometry(table, widths_dxa)
    doc.add_paragraph()
    return table


def set_section(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    header = section.header.paragraphs[0]
    header.text = ""
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_paragraph_spacing(header, after=0)
    r = header.add_run("Duit Cards Backend Architecture")
    set_run_font(r, size=9, color=MUTED)

    footer = section.footer.paragraphs[0]
    footer.text = ""
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_paragraph_spacing(footer, after=0)
    r = footer.add_run("Phase 1 architecture brief")
    set_run_font(r, size=9, color=MUTED)


def configure_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for name in ("List Bullet", "List Number"):
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.5)
        style.paragraph_format.first_line_indent = Inches(-0.25)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.167


def add_title_block(doc):
    add_para(doc, "ARCHITECTURE BRIEF", size=11, bold=True, color=MUTED, after=4)
    add_para(doc, "Duit Cards Backend Architecture - Phase 1", size=23, bold=True, color=BLACK, after=4)
    add_para(
        doc,
        "Backend foundation for digital business cards, relationship context, sharing, onboarding, and future AI/OCR workflows.",
        size=13,
        color=MUTED,
        after=14,
    )

    rows = [
        ("Audience", "Founder/product/engineering team"),
        ("Document status", "Draft architecture baseline"),
        ("Date", "July 13, 2026"),
        ("Current app state", "Expo + React Native frontend prototype with mock data"),
        ("Recommended backend path", "TypeScript API, Postgres, object storage, JWT/session auth"),
    ]
    for label, value in rows:
        p = doc.add_paragraph()
        set_paragraph_spacing(p, after=2)
        r = p.add_run(f"{label}: ")
        set_run_font(r, size=11, color=BLACK, bold=True)
        r2 = p.add_run(value)
        set_run_font(r2, size=11, color=BLACK)

    add_para(doc, "", after=2)
    add_callout(
        doc,
        "Recommendation",
        "Begin with a narrow backend that persists users, onboarding profile, cards, connections, meeting context, and follow-up state. Defer OCR, AI enrichment, CRM integrations, and complex analytics until the core data model is stable.",
    )


def build_doc():
    doc = Document()
    set_section(doc)
    configure_styles(doc)
    add_title_block(doc)

    add_heading(doc, "1. Executive Summary")
    add_para(
        doc,
        "Duit Cards is currently a frontend prototype for digital business cards and lightweight relationship management. The backend should convert mock, in-memory UI behavior into durable user-owned data without overbuilding platform capabilities too early.",
    )
    add_para(
        doc,
        "Phase 1 should focus on the product loop already visible in the app: create a card, share it, save people met in person, retain context, and follow up later. The backend should provide clean app-facing APIs while leaving room for AI, OCR, and CRM sync as extension layers.",
    )

    add_heading(doc, "2. Phase 1 Goals")
    for item in [
        "Persist authenticated user profiles and onboarding answers.",
        "Support one or more digital business cards per user, including public share handles.",
        "Store connections created from card exchanges, manual entry, or future scan/import flows.",
        "Store meeting context, tags, priority placement, notes, and next-step fields.",
        "Expose a typed API contract that the React Native app can consume incrementally.",
        "Keep the architecture simple enough to ship, test, and operate as a small product team.",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "3. Non-Goals For Phase 1")
    for item in [
        "Full business-card OCR pipeline.",
        "Real-time messaging or social networking feed.",
        "Enterprise CRM synchronization.",
        "Complex recommendation systems or agentic AI workflows.",
        "Multi-tenant organization administration.",
        "Payment, subscription, or billing infrastructure.",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "4. Recommended System Architecture")
    add_para(doc, "The recommended baseline is a conventional mobile API architecture:")
    for item in [
        "React Native mobile app calls a versioned HTTPS API.",
        "API service handles authentication, authorization, validation, business rules, and response shaping.",
        "Postgres stores relational product data.",
        "Object storage stores profile photos, uploaded card images, and future scan assets.",
        "Background workers handle asynchronous tasks such as image processing, website enrichment, reminders, and AI calls.",
        "Observability collects structured logs, API metrics, traces, and error reports.",
    ]:
        add_bullet(doc, item)

    add_table(
        doc,
        ["Layer", "Responsibility", "Recommended Choice"],
        [
            ("Mobile client", "Screens, local cache, optimistic UI, secure token storage", "Existing Expo React Native app"),
            ("API", "REST or typed RPC endpoints, validation, auth, domain logic", "Node.js/TypeScript with Express, NestJS, or Fastify"),
            ("Database", "Source of truth for users, cards, connections, reminders", "Postgres"),
            ("ORM/migrations", "Schema evolution and typed data access", "Prisma or Drizzle"),
            ("Object storage", "Images and future scanned card assets", "S3-compatible storage or Supabase Storage"),
            ("Jobs", "AI/OCR/reminders/web enrichment outside request path", "Queue + worker, added after MVP"),
        ],
        [1700, 4300, 3360],
    )

    doc.add_page_break()
    add_heading(doc, "5. Suggested Technology Direction")
    add_para(
        doc,
        "For a fast but durable build, use a TypeScript backend with Postgres. Supabase is a strong shortcut if speed is more important than custom infrastructure; a custom Node.js API with Prisma or Drizzle is better if the product needs tighter control over authorization, API contracts, and future job orchestration.",
    )
    add_table(
        doc,
        ["Option", "Best When", "Trade-Off"],
        [
            ("Supabase-first", "You want auth, Postgres, storage, and admin tooling quickly.", "Less custom control; RLS discipline is important."),
            ("Custom Node API + Postgres", "You want a clean domain/API layer and future workers.", "More setup and ownership."),
            ("Next.js API + Prisma", "You also need a web app/admin surface soon.", "Can blur mobile API and web concerns if not structured carefully."),
        ],
        [2300, 3800, 3260],
    )

    add_heading(doc, "6. Domain Model")
    add_para(
        doc,
        "The backend should model the product concepts already present in the frontend, rather than storing raw UI mock objects as-is.",
    )
    add_table(
        doc,
        ["Entity", "Purpose", "Important Fields"],
        [
            ("User", "Authenticated account owner.", "id, email, phone, name, avatarUrl, createdAt"),
            ("UserProfile", "Onboarding/profile details.", "roleTitle, company, website, intents, aiFollowUpChoice"),
            ("BusinessCard", "A shareable card owned by a user.", "title, subtitle, slug, theme, status, publicUrl"),
            ("Connection", "A person the user met or saved.", "name, role, company, city, source, placement, ownerUserId"),
            ("CardExchange", "How cards were shared/received.", "exchangeType, exchangedAt, location, eventName"),
            ("ConnectionNote", "User-authored notes and context.", "body, createdAt, updatedAt"),
            ("Tag", "Reusable labels for filtering.", "name, color, ownerUserId"),
            ("FollowUp", "Reminder or next action.", "dueAt, status, snoozedUntil, completedAt"),
            ("Asset", "Uploaded photos/card images.", "storageKey, contentType, ownerUserId, purpose"),
        ],
        [1900, 3050, 4410],
    )

    doc.add_page_break()
    add_heading(doc, "7. API Surface")
    add_para(
        doc,
        "Use versioned endpoints from the beginning, for example `/v1`. Keep payloads close to the app-level types but avoid leaking database internals.",
    )
    add_table(
        doc,
        ["Area", "Representative Endpoints", "Notes"],
        [
            ("Auth", "POST /v1/auth/register, POST /v1/auth/login, POST /v1/auth/logout", "May be delegated to Supabase/Auth0/Clerk."),
            ("Me/Profile", "GET /v1/me, PATCH /v1/me/profile", "Stores onboarding answers and profile identity."),
            ("Cards", "GET /v1/cards, POST /v1/cards, PATCH /v1/cards/:id", "Supports multiple cards per user."),
            ("Public card", "GET /v1/public/cards/:slug", "Unauthenticated read path for sharing/QR."),
            ("Connections", "GET /v1/connections, POST /v1/connections, PATCH /v1/connections/:id", "Primary mobile feed API."),
            ("Connection detail", "GET /v1/connections/:id", "Includes notes, tags, exchange context, follow-ups."),
            ("Follow-ups", "GET /v1/follow-ups, POST /v1/follow-ups, PATCH /v1/follow-ups/:id", "Reminder list and completion/snooze."),
            ("Assets", "POST /v1/assets/upload-url", "Pre-signed uploads for images/scans."),
        ],
        [1700, 4300, 3360],
    )

    add_heading(doc, "8. Mobile Integration Strategy")
    for item in [
        "Create an API client module in the mobile app before changing screens.",
        "Map API responses into the existing `Connection`, `BusinessCard`, and `Meeting` app-level types.",
        "Replace `socialRepository` mock functions one by one with API-backed functions or hooks.",
        "Add local loading, empty, and error states to each screen as it becomes API-backed.",
        "Persist auth tokens securely using platform-appropriate secure storage.",
        "Keep the app usable with mock data in development until every endpoint exists.",
    ]:
        add_number(doc, item)

    doc.add_page_break()
    add_heading(doc, "9. Authentication And Authorization")
    add_para(
        doc,
        "Every private object should be scoped by `ownerUserId`. Public card URLs should expose only deliberately public card fields, not private notes, connection history, phone numbers, or onboarding answers.",
    )
    add_table(
        doc,
        ["Resource", "Access Rule"],
        [
            ("User profile", "Only the authenticated owner can read/write private profile fields."),
            ("Business card", "Owner can manage; public slug can read public card snapshot."),
            ("Connection", "Only owner can access."),
            ("Notes/follow-ups", "Only owner can access."),
            ("Assets", "Private by default; public only when attached to public card surface."),
        ],
        [2300, 7060],
    )

    add_heading(doc, "10. AI, OCR, And Website Enrichment Extension Points")
    add_para(
        doc,
        "Do not place AI or OCR directly in the mobile app request path. Store the original asset/input, enqueue background work, and write normalized results back into first-class product tables after review or confidence checks.",
    )
    for item in [
        "OCR card scan: uploaded image -> OCR worker -> extracted draft contact -> user review -> saved connection.",
        "Website enrichment: profile website -> safe server-side fetch -> summarized profile suggestions -> user-approved updates.",
        "Follow-up AI: connection context -> suggested next step -> editable recommendation, not automatic action.",
        "Deduplication: new card/contact -> candidate matching job -> merge suggestion.",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "11. Security And Privacy")
    for item in [
        "Treat onboarding answers, contacts, notes, and phone numbers as sensitive personal data.",
        "Use TLS everywhere and store secrets only in managed secret storage.",
        "Keep public-card payloads intentionally small and separate from private profile data.",
        "Add rate limits for public card reads, auth endpoints, upload URLs, and AI/OCR jobs.",
        "Avoid sending raw HTML or full contact histories to AI providers without redaction, limits, and consent.",
        "Log identifiers and operational metadata, not raw notes or private contact bodies.",
    ]:
        add_bullet(doc, item)

    doc.add_page_break()
    add_heading(doc, "12. Deployment And Environments")
    add_table(
        doc,
        ["Environment", "Purpose", "Notes"],
        [
            ("Local", "Developer API, local Postgres, seed data.", "Use Docker Compose or hosted dev database."),
            ("Staging", "Mobile QA and release validation.", "Mirrors production auth/storage shape."),
            ("Production", "Real users and public card links.", "Backups, monitoring, alerts, and migration discipline required."),
        ],
        [1700, 3600, 4060],
    )

    add_heading(doc, "13. Implementation Roadmap")
    add_table(
        doc,
        ["Phase", "Deliverables", "Exit Criteria"],
        [
            ("0 - Decisions", "Choose stack, hosting, auth provider, API style, schema migration tool.", "Architecture doc accepted and backend repo/thread started."),
            ("1 - Backend skeleton", "API app, health check, env config, database connection, migrations, CI checks.", "Deployable empty service with automated tests."),
            ("2 - Auth/profile", "Register/login/session, `/me`, onboarding profile persistence.", "Mobile can save onboarding to backend."),
            ("3 - Cards", "Card CRUD, public slug read, asset upload flow.", "Share screen uses backend card data."),
            ("4 - Connections", "Connection CRUD, tags, placement, detail endpoint.", "Home/detail screens read/write real data."),
            ("5 - Follow-ups", "Reminder status, due dates, snooze/complete.", "Meetings/follow-up surfaces become persistent."),
            ("6 - Extensions", "OCR/AI jobs, enrichment, dedupe, analytics.", "Queued async workflows with user review gates."),
        ],
        [1600, 4750, 3010],
    )

    add_heading(doc, "14. Open Questions")
    for item in [
        "Will Duit Cards use phone-based auth, email auth, social login, or all three?",
        "Should public card slugs be globally unique forever, or can users change them with redirects?",
        "Should connections be private to one user only, or can two users mutually link an exchange later?",
        "What is the first data source for creating a connection: manual form, QR exchange, imported contact, or card scan?",
        "What compliance/privacy bar is required before storing real contact data?",
        "Will the backend live in this mobile repo, a monorepo, or a separate service repository?",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "15. Suggested Next Step")
    add_para(
        doc,
        "Create a new implementation thread after this architecture brief is accepted. The first backend task should scaffold the chosen API service, database schema, migrations, seed data, and a minimal `/health` plus `/v1/me` path. Keep the frontend thread focused on API integration and mobile behavior.",
    )
    add_heading(doc, "Backend Kickoff Checklist", level=2)
    for item in [
        "Pick the backend stack and hosting target.",
        "Decide whether the backend lives in this repo, a monorepo, or a separate service repo.",
        "Create initial database schema and migration workflow.",
        "Create seeded development data based on the current mock contacts/cards.",
        "Define the first mobile API client contract.",
        "Ship `/health`, `/v1/me`, and onboarding profile persistence before expanding scope.",
    ]:
        add_bullet(doc, item)

    doc.save(OUT)


if __name__ == "__main__":
    build_doc()
    print(OUT)
